"""Extract every preventive-maintenance planning from the workbook into compact JSON.

Columns are found from the header labels (tl, Projets / position presse, N° des machines,
Types des machines, N° de série, S1...), so inserting or moving columns in Excel is fine.

Usage: python parse_plan.py <xlsx> <out.json>
"""
import json
import re
import sys
from collections import Counter

import openpyxl

CODES = set("MTSAHDC")

SHEETS = [
    {"match": "cutting", "svc": "coupe"},
    {"match": "cellul", "svc": "cellule"},
    {"match": "pre-ass", "svc": "preassemblage"},
    {"match": "surmoul", "svc": "surmoulage"},
]
SERVICES = [
    {"id": "coupe", "name": "Coupe", "sheet": "Cutting Planning 2026"},
    {"id": "cellule", "name": "Cellule sertissage", "sheet": "CELLULLE"},
    {"id": "preassemblage", "name": "Pré-assemblage", "sheet": "Pre-assembly planning 2026"},
    {"id": "surmoulage", "name": "Surmoulage", "sheet": "PM surmoulage 2026"},
]
# Same person typed differently across sheets
ALIASES = {
    "Abdelali Moustain": "Abdelali Moustaine",
    "Abdelali": "Abdelali Moustaine",
    "Ayoub Elarrafi": "Ayoub El Arrafi",
    "Zazoul Zouhair": "Zouhair Zazoule",
    "Loukrimi Mohamed Rachid Oumjoud": "Loukrimi Mohamed & Rachid Oumjoud",
}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float):
        v = int(v) if v.is_integer() else repr(v)
    s = str(v).replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def title_name(s):
    return " ".join(w.capitalize() if w != "&" else w for w in s.split(" ")) if s else ""


def person(s):
    n = title_name(clean(s))
    return ALIASES.get(n, n)


def merged_lookup(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = (tl, rng.min_row, rng.min_col)
    return m


def slug(s):
    s = s.upper().replace(" ", "")
    s = re.sub(r"[^A-Z0-9._-]", "-", s)
    return s.strip("-") or "X"


def find_layout(ws):
    """Locate the week columns (S1...) and the label columns from the first header block."""
    for r in range(1, 20):
        for c in range(1, ws.max_column + 1):
            if clean(ws.cell(r, c).value) == "S1":
                wcol, hrow = c, r
                n = 0
                for cc in range(wcol, wcol + 60):
                    v = clean(ws.cell(hrow, cc).value)
                    if re.fullmatch(r"S\d+", v):
                        n = int(v[1:])
                    elif n:
                        break
                cols = {}
                for lr in (hrow - 1, hrow):
                    for cc in range(1, wcol):
                        v = clean(ws.cell(lr, cc).value).lower()
                        if not v:
                            continue
                        if v == "tl":
                            cols.setdefault("tl", cc)
                        elif "types des machines" in v:
                            cols.setdefault("type", cc)
                        elif "série" in v or "serie" in v:
                            cols.setdefault("serie", cc)
                        elif "n° des machines" in v or v == "n°":
                            cols.setdefault("num", cc)
                        elif "projet" in v or "position" in v:
                            cols.setdefault("proj", cc)
                assert {"num", "type", "serie"} <= cols.keys(), f"{ws.title}: header labels not found {cols}"
                # people columns: whatever sits left of the machine number, apart from tl / projet
                ppl = [cc for cc in range(1, cols["num"]) if cc not in (cols.get("tl"), cols.get("proj"))]
                if len(ppl) >= 2:
                    cols["chef"], cols["tech"] = ppl[-2], ppl[-1]
                elif ppl:
                    cols["tech"] = ppl[0]
                return wcol, n, cols
    raise SystemExit(f"{ws.title}: no S1 header found")


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    machines, report = [], []
    seen_ids = Counter()
    section_tl = {}
    for ws in wb.worksheets:
        cfg = next((c for c in SHEETS if c["match"] in ws.title.lower()), None)
        if not cfg:
            report.append(f"SKIPPED sheet {ws.title!r}")
            continue
        merges = merged_lookup(ws)
        wcol, nweeks, cols = find_layout(ws)
        report.append(f"{ws.title!r}: weeks={nweeks} cols={ {k: openpyxl.utils.get_column_letter(v) for k, v in cols.items()} }")
        tl_count = Counter()
        carry = {}
        for r in range(1, ws.max_row + 1):
            num = clean(ws.cell(r, cols["num"]).value)
            if not num or num.lower().startswith("n°"):
                if clean(ws.cell(r, wcol).value) == "S1":
                    carry = {}  # new table header: reset forward-fill
                continue
            typ = clean(ws.cell(r, cols["type"]).value)
            serie = clean(ws.cell(r, cols["serie"]).value)
            if not typ and not serie:
                continue
            rec = {"svc": cfg["svc"], "row": r}
            for k in ("tl", "chef", "tech", "proj"):
                if k not in cols:
                    continue
                c = cols[k]
                v = clean(merges[(r, c)][0]) if (r, c) in merges else clean(ws.cell(r, c).value)
                if v:
                    carry[k] = v
                else:
                    v = carry.get(k, "")
                rec[k] = v
            if rec.get("tl"):
                tl_count[person(rec["tl"])] += 1
            rec["num"], rec["type"], rec["serie"] = num, typ, serie
            weeks, note = [], None
            for i in range(nweeks):
                c = wcol + i
                raw = ws.cell(r, c).value
                if (r, c) in merges:
                    _, mr, mc = merges[(r, c)]
                    if (mr, mc) != (r, c):
                        raw = None if mr == r else raw  # a note merged to the right only counts once
                v = clean(raw)
                if not v:
                    weeks.append(".")
                elif v.upper() in CODES and len(v) == 1:
                    weeks.append(v.upper())
                else:
                    weeks.append(".")
                    if note is None:
                        note = {"text": v, "week": i + 1}
            rec["w"] = "".join(weeks)
            if note:
                rec["note"], rec["noteWeek"] = note["text"], note["week"]
                t = note["text"].lower()
                rec["etat"] = "transfere" if "transf" in t else "rebut" if "scrap" in t else "arret"
                rec["w"] = rec["w"][: note["week"] - 1] + "." * (nweeks - note["week"] + 1)
            base = f"{cfg['svc'][:3]}-{slug(num)}"
            seen_ids[base] += 1
            rec["id"] = base if seen_ids[base] == 1 else f"{base}-{seen_ids[base]}"
            machines.append(rec)
        if tl_count:
            section_tl[cfg["svc"]] = tl_count.most_common(1)[0][0]

    for m in machines:
        for k in ("chef", "tech"):
            if m.get(k):
                m[k] = person(m[k])
        if m.get("proj"):
            m["proj"] = m["proj"].upper()

    data = {
        "year": 2026,
        "source": src.replace("\\", "/").split("/")[-1],
        "services": [dict(s, tl=section_tl[s["id"]]) if s["id"] in section_tl else s for s in SERVICES],
        "machines": [
            {k: m[k] for k in ("id", "svc", "chef", "tech", "proj", "num", "type", "serie", "w", "note", "noteWeek", "etat") if m.get(k)}
            for m in machines
        ],
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))

    print("\n".join(report))
    print("machines per service:", dict(Counter(m["svc"] for m in machines)), "total", len(machines))
    print("section TL:", section_tl)
    print("duplicate machine numbers:", [k for k, n in seen_ids.items() if n > 1])
    print("rows missing tech:", [(m["svc"], m["num"]) for m in machines if not m.get("tech")])
    print("json bytes:", len(json.dumps(data, ensure_ascii=False, separators=(",", ":"))))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
